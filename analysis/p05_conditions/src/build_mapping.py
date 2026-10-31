# -*- coding: utf-8 -*-
"""P5 映射表：conditions.jsonl -> outputs/映射表.xlsx。"""
from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))
from lib.paths import SCENIC_AREAS, analysis_dir, scenic_csv  # noqa: E402

ROOT = analysis_dir("p05_conditions")
OUT = ROOT / "outputs"
XLSX = OUT / "映射表.xlsx"

HEADERS = [
    ("cond_id", "条件ID", 11),
    ("scenic_area", "景区", 12),
    ("condition_type", "条件类型", 10),
    ("condition_text", "条件文本", 24),
    ("suggestion_text", "建议文本", 36),
    ("source_url", "来源URL", 34),
    ("doc_id", "文档ID", 10),
    ("trigger_word", "触发词", 9),
    ("source_sentence", "原句", 50),
    ("char_start", "起始位置", 9),
    ("char_end", "结束位置", 9),
    ("is_duplicate", "是否重复", 9),
    ("review_label", "抽查判定", 10),
]


def load_jsonl(path: Path) -> list[dict]:
    return [json.loads(x) for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]


def load_source_urls() -> dict[str, str]:
    """doc_id -> 原始 source_url（从项目1三景区 CSV 取，conditions.jsonl 不带 URL）。"""
    urls: dict[str, str] = {}
    for area in SCENIC_AREAS:
        with scenic_csv(area).open(encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                urls[row["doc_id"]] = row.get("source_url", "")
    return urls


def main() -> int:
    rows = load_jsonl(OUT / "conditions.jsonl")
    review = {}
    rp = OUT / "manual_review.csv"
    if rp.exists():
        with rp.open(encoding="utf-8-sig", newline="") as f:
            review = {r["cond_id"]: r["label"] for r in csv.DictReader(f)}
    stats = json.loads((OUT / "condition_stats.json").read_text(encoding="utf-8"))
    source_urls = load_source_urls()

    wb = Workbook()
    ws = wb.active
    ws.title = "条件映射表"
    ws.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="1F3B63")
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="微软雅黑", size=9)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for j, (key, title, width) in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=j, value=key)
        c.comment = Comment(f"中文：{title}", "P5")
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[get_column_letter(j)].width = width

    for i, r in enumerate(rows, start=2):
        for j, (key, _, _) in enumerate(HEADERS, start=1):
            if key == "review_label":
                v = review.get(r["cond_id"], "未抽查")
            elif key == "source_url":
                v = source_urls.get(r["doc_id"], "")
            else:
                v = r.get(key, "")
            if isinstance(v, bool):
                v = "是" if v else "否"
            c = ws.cell(row=i, column=j, value=v)
            c.font = body_font
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=key in {"suggestion_text", "source_sentence", "source_url"})
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"

    # 统计页
    ws2 = wb.create_sheet("统计")
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 12
    ws2.column_dimensions["D"].width = 12
    ws2.column_dimensions["E"].width = 12
    ws2.column_dimensions["F"].width = 12
    ws2["A1"] = "项目5 条件-建议抽取统计"
    ws2["A1"].font = Font(name="微软雅黑", size=13, bold=True, color="1F3B63")

    def put(row: int, values: list, bold: bool = False):
        for j, v in enumerate(values, start=1):
            c = ws2.cell(row=row, column=j, value=v)
            c.font = Font(name="微软雅黑", size=10, bold=bold)
            c.alignment = Alignment(vertical="center")

    put(3, ["指标", "数值"], True)
    put(4, ["文档总数", stats["docs_total"]])
    put(5, ["有效文档", stats["docs_valid"]])
    put(6, ["有条件的文档", stats["docs_with_conditions"]])
    put(7, ["文档覆盖率", f"{stats['coverage']:.2%}"])
    put(8, ["条件-建议总数", stats["total"]])
    put(9, ["去重后条数", stats["unique"]])
    put(10, ["重复条数", stats["total"] - stats["unique"]])
    put(11, ["人工抽查", len(review)])
    if review:
        cnt = Counter(review.values())
        put(12, ["抽查-正确", cnt.get("正确", 0)])
        put(13, ["抽查-部分正确", cnt.get("部分正确", 0)])
        put(14, ["抽查-错误", cnt.get("错误", 0)])

    put(17, ["条件类型", "条数"], True)
    r0 = 18
    for k, v in sorted(stats["by_type"].items(), key=lambda kv: -kv[1]):
        put(r0, [k, v]); r0 += 1
    put(r0 + 1, ["景区", "条数"], True)
    r0 += 2
    for k, v in sorted(stats["by_area"].items(), key=lambda kv: -kv[1]):
        put(r0, [k, v]); r0 += 1
    put(r0 + 1, ["抽取规则", "条数"], True)
    r0 += 2
    for k, v in sorted(stats["by_rule"].items(), key=lambda kv: -kv[1]):
        put(r0, ["显式条件" if k == "explicit_if" else "隐式条件", v]); r0 += 1

    # 说明页
    ws3 = wb.create_sheet("说明")
    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 96
    notes = [
        ("表名", "项目5 条件-建议映射表"),
        ("数据范围", "长隆 / 白云山 / 武功山 三景区 clean_content（29 条，28 条有效）"),
        ("字段说明", "condition_text 为条件片段，suggestion_text 为建议片段，二者均为 source_sentence 的子串"),
        ("位置口径", "char_start / char_end 为在 clean_content 中的绝对字符位置，左闭右开"),
        ("去重口径", "is_duplicate=true 表示同类型下条件+建议完全重复或前后缀/近似重复，保留原记录"),
        ("抽查口径", "三景区分层抽取 20 条，分为正确 / 部分正确 / 错误；未抽查标“未抽查”"),
        ("列名对照", "cond_id=条件ID；scenic_area=景区；condition_type=条件类型；condition_text=条件文本；"
                     "suggestion_text=建议文本；source_url=来源URL；doc_id=文档ID；trigger_word=触发词；"
                     "source_sentence=原句；char_start=起始位置；char_end=结束位置；is_duplicate=是否重复；"
                     "review_label=抽查判定"),
        ("生成脚本", "analysis/p05_conditions/src/build_mapping.py"),
    ]
    for i, (k, v) in enumerate(notes, start=1):
        a = ws3.cell(row=i, column=1, value=k)
        b = ws3.cell(row=i, column=2, value=v)
        a.font = Font(name="微软雅黑", size=10, bold=True)
        b.font = Font(name="微软雅黑", size=9)
        b.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(XLSX)
    missing = sum(1 for r in rows if not source_urls.get(r["doc_id"]))
    print(f"wrote {XLSX} rows={len(rows)} missing_source_url={missing}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())